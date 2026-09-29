import asyncio, io, sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from datetime import datetime, timezone
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

GOLD = RGBColor(0x8B, 0x6F, 0x2F)
DARK = RGBColor(0x2B, 0x1B, 0x1B)

def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = DARK if level == 1 else GOLD
    return h

def field_table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=2)
    t.style = "Table Grid"
    t.columns[0].width = Cm(6.5)
    t.columns[1].width = Cm(10.5)
    for i, (label, hint) in enumerate(rows):
        c0 = t.cell(i, 0)
        c0.text = label
        for p in c0.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
        c1 = t.cell(i, 1)
        c1.text = hint
        for p in c1.paragraphs:
            for r in p.runs:
                r.italic = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    doc.add_paragraph()
    return t

def note(doc, text):
    p = doc.add_paragraph(text)
    for r in p.runs:
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = GOLD

doc = Document()

title = doc.add_heading("TEMPLATE QUEST — NOTTURNA Young Blood", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph("Scheda compilabile per la Narrazione — L'Archivio Maledetto (Oracolo AI)")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in sub.runs:
    r.italic = True

heading(doc, "Istruzioni per la compilazione")
doc.add_paragraph(
    "• Compila i campi nella colonna di destra (cancella il testo guida in grigio).\n"
    "• Per aggiungere più PNG, Luoghi, Oggetti o Prove: copia e incolla l'intera tabella della sezione.\n"
    "• Indica SEMPRE la Regione: l'Oracolo mostra i contenuti solo ai giocatori della regione giusta (o a tutti se \"Nazionale\").\n"
    "• Le PAROLE CHIAVE sono fondamentali: l'Oracolo trova PNG, luoghi e prove tramite le parole usate dai giocatori in chat.\n"
    "• Una volta compilato, esporta il file in PDF e caricalo nella Knowledge Base dal pannello Narrazione, scegliendo la regione corretta."
)

heading(doc, "1. SPIEGAZIONE QUEST")
note(doc, "Il quadro generale della quest: l'Oracolo lo usa per orientare la narrazione e capire cosa può rivelare e quando.")
field_table(doc, [
    ("Titolo della quest", "es. Il Reliquiario Perduto"),
    ("Regione", "Lazio / Umbria / Abruzzo / Lombardia / Nazionale"),
    ("Periodo di validità", "es. trimestre dicembre-febbraio, oppure 'sempre attiva'"),
    ("Sinossi (2-3 frasi)", "Riassunto della trama visibile ai giocatori"),
    ("Antefatto segreto", "Ciò che è realmente accaduto e che i PG NON sanno all'inizio"),
    ("Obiettivo dei PG", "Cosa devono scoprire/ottenere/impedire"),
    ("Fasi della quest", "1) ... 2) ... 3) ... (tappe principali in ordine)"),
    ("Come si attiva", "Dove o da chi i PG ricevono il primo aggancio (luogo, PNG, voce di corridoio)"),
    ("Condizioni di successo", "Cosa deve accadere perché la quest sia completata"),
    ("Condizioni di fallimento", "Cosa la fa fallire o complica (tempo, morti, Masquerade infranta...)"),
    ("Ricompense", "Informazioni, oggetti, RISORSE, favori, progressi di conoscenza"),
    ("Collegamenti", "Altre quest, PNG o eventi della cronaca collegati"),
    ("Note per l'Oracolo", "Cosa può improvvisare liberamente e cosa NON deve mai rivelare/inventare"),
])

heading(doc, "2. PNG (Personaggio Non Giocante)")
note(doc, "Copia questa tabella per ogni PNG della quest.")
field_table(doc, [
    ("Nome e carica/titolo", "es. Padre Anselmo, custode dell'archivio diocesano"),
    ("Clan e Generazione", "es. Nosferatu, 11ª — oppure 'mortale'"),
    ("Regione e zona", "es. Lazio — Roma, Trastevere"),
    ("Aspetto e tratti distintivi", "Come appare, come parla, tic e manie"),
    ("Carattere e mood iniziale", "ostile / diffidente / neutrale / amichevole"),
    ("Dove si incontra", "Luogo abituale e orari/condizioni"),
    ("Alias e soprannomi", "Altri nomi con cui i giocatori potrebbero cercarlo"),
    ("Info LIVELLO 1 (base)", "Ciò che dice a chiunque con approccio neutro"),
    ("Info LIVELLO 2", "Ciò che rivela dopo una prova superata o buona diplomazia"),
    ("Info LIVELLO 3 (segreto)", "Il segreto che custodisce e la condizione esatta per rivelarlo"),
    ("Prove associate", "es. Carisma + Diplomazia diff. 6 per il Livello 2"),
    ("Reazione ad aggressioni/Discipline", "Come risponde a intimidazioni o poteri usati su di lui"),
    ("Ruolo nella quest", "Cosa sblocca o complica per i PG"),
])

heading(doc, "3. LUOGO")
note(doc, "Copia questa tabella per ogni luogo della quest.")
field_table(doc, [
    ("Nome del luogo", "es. Cripta di San Callisto"),
    ("Regione e zona", "es. Lazio — Roma Sud, Appia Antica"),
    ("Descrizione e atmosfera", "Cosa vede/sente il PG quando entra (2-4 frasi evocative)"),
    ("Chi lo frequenta", "PNG, fazioni, mortali presenti"),
    ("Accesso", "libero / sorvegliato / serve una prova o un invito (specificare)"),
    ("Cosa si può scoprire", "Indizi e informazioni ottenibili qui"),
    ("Oggetti presenti", "Oggetti trovabili (rimando alla sezione OGGETTI)"),
    ("Eventi ricorrenti", "Cosa accade regolarmente (riti, incontri, ronde)"),
    ("Pericoli", "Minacce per chi esplora (guardiani, trappole, Masquerade)"),
    ("Segreti del luogo", "Cosa è nascosto e la condizione/prova per scoprirlo"),
])

heading(doc, "4. OGGETTO")
note(doc, "Copia questa tabella per ogni oggetto della quest.")
field_table(doc, [
    ("Nome dell'oggetto", "es. Diario cifrato del Siniscalco"),
    ("Descrizione", "Aspetto e natura dell'oggetto"),
    ("Dove si trova / chi lo possiede", "Luogo o PNG che lo custodisce"),
    ("Costo in RISORSE", "0 = gratuito; altrimenti indicare il valore"),
    ("Requisiti per ottenerlo", "Prova, scambio, favore, furto... (specificare)"),
    ("Effetto / utilità", "Cosa permette di fare o scoprire"),
    ("Ruolo nella quest", "In quale fase serve e cosa sblocca"),
    ("Note", "Limitazioni, usi unici, rischi"),
])

heading(doc, "5. PROVA CONTRAPPOSTA")
note(doc, "Copia questa tabella per ogni prova. I punteggi dei PG sono calcolati automaticamente dalla scheda ufficiale.")
field_table(doc, [
    ("Nome della prova", "es. Decifrare il diario"),
    ("Parole chiave di attivazione", "parole che il giocatore userà in chat, separate da virgola (es. diario, cifrato, decifrare)"),
    ("Attributo + Abilità", "es. Intelligenza + Occulto (usare i nomi esatti della scheda)"),
    ("Difficoltà", "valore numerico (es. 6)"),
    ("Categoria di conoscenza", "per il conteggio dei pallini (es. Occulto, Politica, Mondo Oscuro...)"),
    ("Successo", "Cosa scopre/ottiene il PG se supera la prova"),
    ("Fallimento", "Cosa accade se fallisce (falsa pista, allarme, nulla)"),
    ("Pareggio", "Come gestire il pareggio (di norma conta come successo parziale)"),
    ("Collegata a", "PNG / Luogo / Oggetto / fase della quest di riferimento"),
    ("Note", "Ripetibile? (di norma NO) Altre condizioni"),
])

doc.add_paragraph()
footer = doc.add_paragraph("NOTTURNA Young Blood — Template quest per L'Archivio Maledetto. Compilare, esportare in PDF e caricare nella Knowledge Base.")
for r in footer.runs:
    r.italic = True
    r.font.size = Pt(8)

buf = io.BytesIO()
doc.save(buf)
data = buf.getvalue()

async def store():
    from core import db
    await db.upload_files.update_one(
        {"filename": "TEMPLATE-QUEST-NOTTURNA.docx"},
        {"$set": {
            "filename": "TEMPLATE-QUEST-NOTTURNA.docx",
            "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "data": data,
            "created_at": datetime.now(timezone.utc).isoformat(),
        }},
        upsert=True,
    )
    print(f"salvato in MongoDB: {len(data)} bytes")

asyncio.run(store())
